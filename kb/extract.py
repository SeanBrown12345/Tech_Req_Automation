"""Turn worksheet rows into normalized requirement records using a profile."""

import re
from collections import Counter
from dataclasses import dataclass, field

from openpyxl import load_workbook

from kb.profile import Profile, SheetProfile

# Cell contents that mean "nothing here" (template dashes, redaction runs, blanks).
_PLACEHOLDER = re.compile(r"^(-+|x{4,}|\s*)$", re.IGNORECASE)
# List-item children such as "a) Land", "(b) Buildings", "1. Water".
_LIST_ITEM = re.compile(r"^\(?[a-z0-9]{1,2}[).]\s", re.IGNORECASE)
# Requirements that introduce a list of children on the following rows.
_LEAD_IN = re.compile(r"(:\s*$|all that apply|the following\s*$)", re.IGNORECASE)
# Hierarchical IDs: "4.1" -> "4", "5.1a" -> "5.1", "10.2d" -> "10.2".
_CHILD_ID = re.compile(r"^(?P<parent>.+?)(\.\d+|[a-z])$", re.IGNORECASE)


def clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return None if _PLACEHOLDER.match(text) else text


@dataclass
class Record:
    record_id: str
    sheet: str
    row_num: int
    module: str | None
    req_id: str | None
    category: str | None
    subcategory: str | None
    section: str | None
    parent_text: str | None
    requirement: str
    status: str | None
    answer_raw: dict[str, str]
    comment: str | None
    attributes: dict[str, str]
    warnings: list[str] = field(default_factory=list)


@dataclass
class SheetReport:
    sheet: str
    records: int = 0
    skipped_unanswered: int = 0
    statuses: Counter = field(default_factory=Counter)
    unknown_values: Counter = field(default_factory=Counter)  # (signal, raw value) -> count
    multi_marked: int = 0


def _evaluate(sheet: SheetProfile, row: tuple, header: tuple, report: SheetReport):
    """Apply answer signals in order. Returns (status, raw answers, warnings)."""
    status, raw, warnings = None, {}, []
    for signal in sheet.signals:
        can_set = signal.refines is None or status is None or status in signal.refines
        if signal.marks is not None:
            marked = [c for c in signal.marks if c < len(row) and clean(row[c])]
            if not marked:
                continue
            if len(marked) > 1:
                report.multi_marked += 1
                warnings.append(f"{signal.name}: multiple columns marked")
            col = marked[0]
            label = clean(header[col]) if col < len(header) else None
            raw[signal.name] = label or clean(row[col])
            if can_set:
                status = signal.marks[col]
        else:
            value = clean(row[signal.column]) if signal.column < len(row) else None
            if value is None:
                continue
            raw[signal.name] = value
            known, mapped = signal.scale.lookup(value)
            if not known:
                report.unknown_values[(signal.name, value)] += 1
                warnings.append(f"{signal.name}: unmapped value {value!r}")
            elif mapped is not None and can_set:
                status = mapped
    return status, raw, warnings


def iter_requirements(rows, first_row: int, cols: dict[str, int], header_requirement: str | None):
    """Yield (row_num, row, req_id, requirement, section, parent_text) for each requirement row.

    Tracks the context a row only makes sense in: the current section heading, and the parent
    requirement for list children ("a) Land" under "...the following categories of assets:",
    "5.1a" under "5.1").
    """
    def cell(row, role):
        idx = cols.get(role)
        return clean(row[idx]) if idx is not None and idx < len(row) else None

    section = None
    text_by_id: dict[str, str] = {}  # req_id -> requirement text, for hierarchical parents
    lead_in: str | None = None

    for row_num, row in enumerate(rows, start=first_row):
        requirement = cell(row, "requirement")
        req_id = cell(row, "req_id")

        if not requirement:
            # Heading row ("INTEGRATION & DATA MANAGEMENT", "SaaS 2.0 | Security | - | -").
            texts = [t for t in (clean(v) for v in row) if t]
            if texts:
                section = max(texts, key=len)
                lead_in = None
            continue
        if requirement == header_requirement:  # header repeated mid-sheet
            continue

        if req_id:
            text_by_id[req_id] = requirement

        parent = None
        m = _CHILD_ID.match(req_id or "")
        if m and m.group("parent") in text_by_id:
            parent = text_by_id[m.group("parent")]
        elif lead_in and (_LIST_ITEM.match(requirement) or len(requirement.split()) <= 6):
            parent = lead_in
        else:
            lead_in = None
        if _LEAD_IN.search(requirement):
            lead_in = requirement

        yield row_num, row, req_id, requirement, section, parent


def extract_sheet(profile: Profile, sheet: SheetProfile, ws) -> tuple[list[Record], SheetReport]:
    report = SheetReport(sheet.name)
    records: list[Record] = []

    header = next(ws.iter_rows(min_row=sheet.header_row, max_row=sheet.header_row, values_only=True), ())
    cols = sheet.columns
    header_requirement = clean(header[cols["requirement"]]) if cols["requirement"] < len(header) else None
    if not header_requirement:
        raise ValueError(f"{profile.path}: sheet {sheet.name!r} row {sheet.header_row} has no header "
                         "in the requirement column - check header_row / columns")

    def cell(row, role):
        idx = cols.get(role)
        return clean(row[idx]) if idx is not None and idx < len(row) else None

    for row_num, row, req_id, requirement, section, parent in iter_requirements(
            ws.iter_rows(min_row=sheet.first_data_row, values_only=True), sheet.first_data_row, cols,
            header_requirement):
        status, raw, warnings = _evaluate(sheet, row, header, report)
        comment = cell(row, "comment")
        if status is None and comment and not warnings:
            status = "NARRATIVE"

        if not raw and not comment:
            report.skipped_unanswered += 1
            continue

        attributes = {}
        for name, idx in sheet.attributes.items():
            value = clean(row[idx]) if idx < len(row) else None
            if value:
                attributes[name] = value

        records.append(Record(
            record_id=f"{profile.source_id}:{sheet.name}:{row_num}",
            sheet=sheet.name,
            row_num=row_num,
            module=sheet.module,
            req_id=req_id,
            category=cell(row, "category"),
            subcategory=cell(row, "subcategory"),
            section=section,
            parent_text=parent,
            requirement=requirement,
            status=status,
            answer_raw=raw,
            comment=comment,
            attributes=attributes,
            warnings=warnings,
        ))
        report.records += 1
        report.statuses[status] += 1

    return records, report


def extract_profile(profile: Profile) -> tuple[list[Record], list[SheetReport]]:
    wb = load_workbook(profile.file, read_only=True, data_only=True)
    records, reports = [], []
    try:
        for sheet in profile.sheets:
            if sheet.name not in wb.sheetnames:
                raise KeyError(f"{profile.path}: sheet {sheet.name!r} not in workbook (has {wb.sheetnames})")
            recs, report = extract_sheet(profile, sheet, wb[sheet.name])
            records.extend(recs)
            reports.append(report)
    finally:
        wb.close()
    return records, reports
