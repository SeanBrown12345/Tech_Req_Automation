import hashlib
import io
import re
import tempfile
from pathlib import Path

import pandas as pd
import streamlit as st
import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.common import STATUS_LABELS, connect, embedder, status_help, status_label
from kb import settings, store
from kb.detect import distinct_values, guess_sheet, guess_status, read_rows
from kb.extract import clean, extract_profile
from kb.profile import load_scales, parse_profile

NO_INFO = "— no information —"
LABEL_TO_STATUS = {label: code for code, label in STATUS_LABELS.items()}
STATUS_OPTIONS = list(STATUS_LABELS.values())
ROLES = [  # (role, label, required)
    ("requirement", "Requirement text", True),
    ("req_id", "Requirement ID", False),
    ("category", "Category / functional group", False),
    ("subcategory", "Sub-category / topic", False),
    ("comment", "Our comment / narrative", False),
]


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (text or "").lower()).strip("_")[:60]


@st.cache_data(show_spinner="Reading workbook...", max_entries=5)
def analyze(data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = {}
        for ws in wb.worksheets:
            rows = read_rows(ws)
            sheets[ws.title] = (rows, guess_sheet(ws.title, rows))
        return sheets
    finally:
        wb.close()


def repo_source_ids() -> set[str]:
    return {p.stem for p in settings.REPO_PROFILES_DIR.glob("*.yaml")}


def data_profiles() -> list[Path]:
    """Worksheets added from the dashboard, latest upload first (the profile is written on each save)."""
    return sorted(settings.DATA_PROFILES_DIR.glob("*.yaml"), key=lambda p: p.stat().st_mtime, reverse=True)


def sheet_editor(fhash: str, name: str, rows: list[tuple], guess) -> dict | None:
    """Widgets for one sheet. Returns its profile spec, or None if excluded/incomplete."""
    k = f"{fhash}:{name}"
    include = st.toggle("Load this sheet", value=bool(guess.answer_columns or guess.mark_columns), key=f"{k}:inc")
    if not include:
        st.caption("Skipped.")
        return None

    c1, c2 = st.columns([1, 3])
    header_row = c1.number_input("Header row", 1, max(len(rows), 1), guess.header_row, key=f"{k}:hdr")
    module = c2.text_input("Module (optional)", key=f"{k}:mod", placeholder="e.g. CIS, ERP, MDM")

    header = rows[header_row - 1] if header_row <= len(rows) else ()
    letters = [get_column_letter(i + 1) for i in range(len(header))]
    headers = {l: str(v).strip() for l, v in zip(letters, header) if clean(v)}
    fmt = lambda l: f"{l} · {headers[l][:50]}" if l and l in headers else (l or "—")
    choices = list(headers)

    st.markdown("**Columns**")
    columns = {}
    cols = st.columns(len(ROLES))
    for (role, label, required), col in zip(ROLES, cols):
        guessed = guess.columns.get(role)
        options = choices if required else [None, *choices]
        index = options.index(guessed) if guessed in options else 0
        value = col.selectbox(label, options, index=index, format_func=fmt, key=f"{k}:{role}:{header_row}")
        if value:
            columns[role] = value

    st.markdown("**How the answer is recorded**")
    a1, a2 = st.columns(2)
    used = set(columns.values())
    marks = a1.multiselect(
        "Marked columns (an X or letter in one column picks the answer)", [l for l in choices if l not in used],
        default=[l for l in guess.mark_columns if l in choices and l not in used], format_func=fmt,
        key=f"{k}:marks:{header_row}")
    answers = a2.multiselect(
        "Answer value columns (a dropdown or code like Y / N)", [l for l in choices if l not in used | set(marks)],
        default=[l for l in guess.answer_columns if l in choices and l not in used | set(marks)], format_func=fmt,
        key=f"{k}:ans:{header_row}")
    if not marks and not answers:
        st.warning("Pick at least one marked column or answer value column.")
        return None

    signals = []
    st.markdown("**Map their answers to our statuses**", help=status_help())
    if marks:
        df = pd.DataFrame({"Column": marks, "Header": [headers[l] for l in marks],
                           "Status": [status_label(guess_status(headers[l], guess.legend)) if guess_status(
                               headers[l], guess.legend) else NO_INFO for l in marks]})
        edited = st.data_editor(df, hide_index=True, disabled=["Column", "Header"], key=f"{k}:mmap:{header_row}",
                                column_config={"Status": st.column_config.SelectboxColumn(
                                    options=[*STATUS_OPTIONS, NO_INFO], required=True)})
        mapping = {r.Column: LABEL_TO_STATUS[r.Status] for r in edited.itertuples() if r.Status != NO_INFO}
        if mapping:
            signals.append({"name": "marks", "marks": mapping})

    for position, letter in enumerate(answers):
        secondary = bool(signals) or position > 0
        counts = distinct_values(rows, letter, header_row + 1, columns.get("requirement"))
        suggested = []
        for value in counts:
            status = guess_status(value, guess.legend)
            if secondary and status in ("SUPPORTED", "NOT_APPLICABLE"):
                status = None  # a follow-up column saying "supported" adds nothing to the main answer
            suggested.append(status_label(status) if status else NO_INFO)
        df = pd.DataFrame({"Value": list(counts), "Rows": list(counts.values()), "Status": suggested})
        st.caption(f"{fmt(letter)}" + (" — refines the main answer (only sharpens a positive one)" if secondary else ""))
        edited = st.data_editor(df, hide_index=True, disabled=["Value", "Rows"], key=f"{k}:vmap:{letter}:{header_row}",
                                column_config={"Status": st.column_config.SelectboxColumn(
                                    options=[*STATUS_OPTIONS, NO_INFO], required=True)})
        spec = {"name": slug(headers[letter]) or f"col_{letter}", "column": letter,
                "values": {r.Value: (None if r.Status == NO_INFO else LABEL_TO_STATUS[r.Status])
                           for r in edited.itertuples()}}
        if secondary:
            spec["refines"] = ["SUPPORTED", "STANDARD"]
        signals.append(spec)

    leftover = [l for l in choices if l not in used | set(marks) | set(answers)]
    keep = st.multiselect("Other columns to keep as details", leftover, format_func=fmt, key=f"{k}:attrs:{header_row}")

    spec = {"name": name, "header_row": int(header_row), "columns": columns, "answer": signals}
    if module.strip():
        spec["module"] = module.strip()
    if keep:
        spec["attributes"] = {slug(headers[l]) or f"col_{l}": l for l in keep}
    return spec


def build_profile(file_ref: str, source: dict, sheet_specs: list[dict]) -> dict:
    return {"file": file_ref, "source": source, "sheets": sheet_specs}


# =============================================================================================
st.title("Add a completed worksheet")
st.caption("Upload a worksheet we've already submitted. Confirm how it's laid out, preview, then load it "
           "into the knowledge base.")

if loaded_msg := st.session_state.pop("add_ws:loaded", None):
    st.success(loaded_msg)

# Bumping the uploader's key after a save clears it, so the page is ready for the next worksheet.
upload_round = st.session_state.setdefault("add_ws:round", 0)
uploaded = st.file_uploader("Completed worksheet (.xlsx / .xlsm)", type=["xlsx", "xlsm"],
                            key=f"add_ws:upload:{upload_round}")

if uploaded:
    data = uploaded.getvalue()
    fhash = hashlib.sha256(data).hexdigest()[:12]
    sheets = analyze(data)

    st.subheader("1. About this worksheet")
    d1, d2 = st.columns(2)
    client = d1.text_input("Client *", key=f"{fhash}:client", placeholder="e.g. City of Dayton")
    rfp = d2.text_input("RFP / project", key=f"{fhash}:rfp", placeholder="e.g. CIS Replacement")
    worksheet = d1.text_input("Worksheet name", value=Path(uploaded.name).stem, key=f"{fhash}:ws")
    submitted = d2.date_input("Submitted on *", value=None, key=f"{fhash}:date", format="YYYY-MM-DD")
    products = st.text_input("Products proposed (comma-separated)", key=f"{fhash}:prod")
    source_id = slug(f"{client}_{worksheet}") if client else ""
    if source_id in repo_source_ids():
        st.error(f"ID `{source_id}` is already used by a worksheet managed in git. Change the worksheet name.")
        source_id = ""
    elif source_id and (settings.DATA_PROFILES_DIR / f"{source_id}.yaml").exists():
        st.warning(f"A worksheet with ID `{source_id}` was already added; loading will replace it.")

    st.subheader("2. Layout of each sheet")
    found = {n: v for n, v in sheets.items() if v[1] is not None}
    skipped = [n for n, v in sheets.items() if v[1] is None]
    if skipped:
        st.caption("No requirements table detected on: " + ", ".join(skipped))
    if not found:
        st.error("Couldn't find a header row on any sheet.")
        st.stop()

    specs = []
    for tab, (name, (rows, guess)) in zip(st.tabs(list(found)), found.items()):
        with tab:
            spec = sheet_editor(fhash, name, rows, guess)
            if spec:
                specs.append(spec)

    st.subheader("3. Preview and load")
    source = {"client": client.strip(), "rfp": rfp.strip() or None, "worksheet": worksheet.strip() or None,
              "submitted": submitted.isoformat() if submitted else None,
              "products": [p.strip() for p in products.split(",") if p.strip()]}
    problems = [msg for ok, msg in ((client.strip(), "Enter the client."), (submitted, "Enter the submitted date."),
                                    (specs, "Choose at least one sheet to load."), (source_id, "Fix the ID above."))
                if not ok]
    for msg in problems:
        st.caption(f"• {msg}")

    p1, p2 = st.columns(2)
    preview = p1.button("Preview", disabled=not specs, width="stretch")
    save = p2.button("Save and load into knowledge base", type="primary", disabled=bool(problems),
                     width="stretch")
    suffix = Path(uploaded.name).suffix.lower()
    scales = load_scales(settings.SCALES_FILE)

    if preview:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / f"preview{suffix}"
            path.write_bytes(data)
            try:
                profile = parse_profile(build_profile(path.name, source, specs), source_id or "preview", scales, Path(tmp))
                records, reports = extract_profile(profile)
            except Exception as exc:
                st.error(f"Couldn't read the worksheet with these settings: {exc}")
                st.stop()
        for rep in reports:
            with st.container(border=True):
                st.markdown(f"**{rep.sheet}** — {rep.records:,} requirements, {rep.skipped_unanswered} unanswered rows skipped")
                st.write(", ".join(f"{status_label(s)}: {n:,}" for s, n in rep.statuses.most_common()))
                for (signal, value), n in rep.unknown_values.most_common():
                    st.warning(f"Unmapped value {value!r} in {signal} ({n} rows)")
                if rep.multi_marked:
                    st.warning(f"{rep.multi_marked} rows have more than one column marked; the first is used.")
        st.dataframe(pd.DataFrame([{
            "Sheet": r.sheet, "Row": r.row_num, "ID": r.req_id, "Under": r.parent_text, "Requirement": r.requirement,
            "Status": status_label(r.status), "Comment": r.comment} for r in records[:300]]),
            hide_index=True, width="stretch")
        if len(records) > 300:
            st.caption(f"Showing the first 300 of {len(records):,}.")

    if save:
        settings.UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
        settings.DATA_PROFILES_DIR.mkdir(parents=True, exist_ok=True)
        workbook_path = settings.UPLOADS_DIR / f"{source_id}{suffix}"
        workbook_path.write_bytes(data)
        raw = build_profile(workbook_path.relative_to(settings.DATA_DIR).as_posix(), source, specs)
        profile_path = settings.DATA_PROFILES_DIR / f"{source_id}.yaml"
        profile_path.write_text(yaml.safe_dump(raw, sort_keys=False, allow_unicode=True), encoding="utf-8")
        with st.spinner("Loading and indexing..."):
            conn = connect()
            profile = parse_profile(raw, source_id, scales, settings.DATA_DIR, profile_path)
            records, reports = extract_profile(profile)
            store.replace_source(conn, profile, records)
            added = store.embed_missing(conn, embedder())
        st.session_state["add_ws:loaded"] = (
            f"Loaded **{len(records):,}** requirements from {uploaded.name} ({added:,} new embeddings). "
            "They're searchable now.")
        st.session_state["add_ws:round"] = upload_round + 1
        st.rerun()

# ---- Manage worksheets added here ------------------------------------------------------------
st.divider()
st.subheader("Uploaded Worksheets")
added_profiles = data_profiles()
if not added_profiles:
    st.caption("None yet. Worksheets defined in the repository (config/profiles) are managed in git.")
for path in added_profiles:
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    src = raw.get("source") or {}
    row = st.columns([4, 2, 1])
    row[0].write(f"**{src.get('client')}** — {src.get('worksheet') or path.stem}")
    row[1].caption(f"Submitted {src.get('submitted') or '?'} · `{path.stem}`")
    with row[2].popover("Remove"):
        st.write("Remove this worksheet and its answers from the knowledge base?")
        if st.button("Remove", key=f"rm:{path.stem}", type="primary"):
            store.delete_source(connect(), path.stem)
            (settings.DATA_DIR / raw["file"]).unlink(missing_ok=True)
            path.unlink()
            st.rerun()
